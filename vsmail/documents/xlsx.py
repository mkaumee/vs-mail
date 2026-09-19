"""Spreadsheet attachments — 22 files across 15 emails.

The brief never mentions Excel; the bundle uses it anyway, both as the SI
beside a Word BL and as both halves of a pair.
"""
from __future__ import annotations

import io

import openpyxl

from vsmail.documents.base import DocumentUnreadable
from vsmail.documents.tabular import flatten, render_row


def read_xlsx(data: bytes) -> str:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentUnreadable(f"spreadsheet could not be opened: {exc}") from exc

    lines: list[str] = []
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            if len(workbook.sheetnames) > 1:
                lines.append(f"[sheet: {name}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [flatten(str(cell)) for cell in row if cell not in (None, "")]
                line = render_row(cells)
                if line:
                    lines.append(line)
    finally:
        workbook.close()

    text = "\n".join(lines).strip()
    if not text:
        raise DocumentUnreadable("spreadsheet contains no readable cells")
    return text
